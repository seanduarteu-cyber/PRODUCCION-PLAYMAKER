# Receta de producción Playmaker: un solo archivo, 4 hojas (Ficha, Moldes, Despiece, Precios).
# Toma los datos ya cargados en v6 (telas, confección, avíos, moldes, impresión).
import sys, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

SRC, DST = sys.argv[1], sys.argv[2]
import json
TELAS = json.load(open(sys.argv[3]))   # listado de telas actualizado
v6 = load_workbook(SRC)
AZUL = PatternFill('solid', fgColor='DDEBF7'); GRIS = PatternFill('solid', fgColor='F2F2F2')
HF = PatternFill('solid', fgColor='1F3864'); AMA = PatternFill('solid', fgColor='FFF2CC'); NAR = PatternFill('solid', fgColor='FCE4D6')
H = Font(bold=True, color='FFFFFF'); B = Font(bold=True); ROJO = Font(color='C00000', size=9)
MONEY = '"$"#,##0;-"$"#,##0;"$"0'; PCT = '0.0%'; M3 = '0.000'
WR = Alignment(wrap_text=True, vertical='top'); WC = Alignment(wrap_text=True, vertical='center', horizontal='center')
LINEA = Border(top=Side(style='medium', color='1F3864'))

wb = Workbook()
def titulo(ws, t, sub):
    ws['A1'] = t; ws['A1'].font = Font(bold=True, size=14, color='1F3864')
    ws['A2'] = sub; ws['A2'].font = Font(italic=True, color='595959')
def enc(ws, fila, col0, cols, anchos):
    for i, (t, w) in enumerate(zip(cols, anchos)):
        c = ws.cell(fila, col0 + i, t); c.font = H; c.fill = HF; c.alignment = WC
        if w: ws.column_dimensions[L(col0 + i)].width = w
    ws.row_dimensions[fila].height = 42
def lista(ws, formula, rango):
    d = DataValidation(type='list', formula1=formula, allow_blank=True); ws.add_data_validation(d); d.add(rango)

# ---------------------------------------------------------------- datos del v6
telas6, avios6 = [], []
for r in v6['Precios'].iter_rows(min_row=5, max_row=120, max_col=8, values_only=True):
    if r[0]: (telas6 if r[1] == 'Tela' else avios6).append(r)
conf6 = [r for r in v6['Confeccion'].iter_rows(min_row=5, max_row=100, max_col=2, values_only=True) if r[0]]
rec6 = [r for r in v6['Recetas'].iter_rows(min_row=5, max_row=84, max_col=17, values_only=True) if r[0] and not str(r[0]).startswith('PRENDAS SIN')]
# Moldes del catálogo vigente (CATALOGO_MOLDES_VIGENTES_INDEXADO_2.pdf) que no estaban: 55-65 y especiales E01-E09.
# Piezas sugeridas según el dibujo de cada guía de tallas: revisar en la marcada.
SM_, RG, CC, TAP, CU = ('Ribete cuello y sisas', 1), ('Manga raglán', 2), ('Cuello camisero', 1), ('Tapeta', 1), ('Tira de cuello', 1)
DE, ES, MA, PL = ('Delantero', 1), ('Espalda', 1), ('Manga', 2), ('Paneles laterales', 2)
NUEVOS = [  # código, descripción, deporte, prenda, calce, confección, piezas
 ('55 · VOMUJV03', 'Vóley, sin mangas, cuello v', 'Vóley', 'Camiseta', 'Femenino', 'Camiseta Voley', [DE, ES, SM_]),
 ('56 · VOMUJMAV03', 'Vóley, pro, sin mangas, cuello v', 'Vóley', 'Camiseta', 'Femenino', 'Camiseta Voley', [DE, ES, SM_]),
 ('57 · VOMUJVCC03', 'Vóley, pro, sin mangas, cuello v con cuello', 'Vóley', 'Camiseta', 'Femenino', 'Camiseta Voley', [DE, ES, SM_, CC]),
 ('58 · RB-HOMU02', 'Rugby, raglán, cuello redondo', 'Rugby', 'Camiseta', 'Masculino', 'Camiseta De Rugby Básica', [DE, ES, RG, CU]),
 ('59 · RB-MUJU02', 'Rugby, raglán, cuello redondo', 'Rugby', 'Camiseta', 'Femenino', 'Camiseta De Rugby Básica', [DE, ES, RG, CU]),
 ('60 · RBHOMCC02', 'Rugby, raglán, cuello camisero', 'Rugby', 'Camiseta', 'Masculino', 'Camiseta Rugby Con Cuello', [DE, ES, RG, CC, TAP]),
 ('61 · RBMUJCC02', 'Rugby, raglán, cuello camisero, cortes laterales', 'Rugby', 'Camiseta', 'Femenino', 'Camiseta Rugby Con Cuello', [DE, ES, RG, PL, CC, TAP]),
 ('62 · RB-HOMCL02', 'Rugby, clásica, cuello camisero', 'Rugby', 'Camiseta', 'Masculino', 'Camiseta Rugby Con Cuello', [DE, ES, MA, CC, TAP]),
 ('63 · RBMUJCL02', 'Rugby, clásica, cuello camisero', 'Rugby', 'Camiseta', 'Femenino', 'Camiseta Rugby Con Cuello', [DE, ES, MA, CC, TAP]),
 ('64 · HO-MUJ03', 'Hockey, sin mangas, cuello redondo', 'Hockey', 'Camiseta', 'Femenino', 'Camiseta De Hockey', [DE, ES, SM_]),
 ('65 · HO-MUJCC03', 'Hockey, sin mangas, cuello camisero', 'Hockey', 'Camiseta', 'Femenino', 'Camiseta De Hockey', [DE, ES, SM_, CC, TAP]),
 ('E01 · RB-HOMUCC02', 'Especial rugby, manga corta, cortes laterales', 'Rugby', 'Camiseta', 'Masculino', 'Camiseta De Rugby Básica', [DE, ES, MA, PL, CU]),
 ('E02 · RB-HOMUCC02 SM', 'Especial rugby, sin mangas, cortes laterales', 'Rugby', 'Camiseta', 'Masculino', 'Camiseta De Rugby Básica', [DE, ES, PL, SM_]),
 ('E03 · RBHOMMA02', 'Especial rugby, raglán, cuello con tapeta', 'Rugby', 'Camiseta', 'Masculino', 'Camiseta Rugby Con Cuello', [DE, ES, RG, CU, TAP]),
 ('E04 · FUTHOMMA01', 'Especial fútbol, cuello con tapeta', 'Fútbol', 'Camiseta', 'Masculino', 'Camiseta De Futbol', [DE, ES, MA, CU, TAP]),
 ('E05 · CA-MUJ', 'Calza corta', 'Varios', 'Calza', 'Femenino', 'Calza', [('Pierna', 2), ('Pretina', 1)]),
 ('E06 · C-FALDA TUBO', 'Calza-falda tubo', 'Varios', 'Falda', 'Femenino', 'Calza-Falda', [('Falda delantero', 1), ('Falda trasero', 1), ('Calza pierna', 2), ('Pretina', 1)]),
 ('E07 · C-FALDA CRUZADA', 'Calza-falda cruzada', 'Varios', 'Falda', 'Femenino', 'Calza-Falda', [('Falda delantero cruzado', 2), ('Falda trasero', 1), ('Calza pierna', 2), ('Pretina', 1)]),
 ('E08 · C-FALDA CON COSTADO', 'Calza-falda con costado', 'Varios', 'Falda', 'Femenino', 'Calza-Falda', [('Falda delantero', 1), ('Falda trasero', 1), ('Costados', 2), ('Calza pierna', 2), ('Pretina', 1)]),
 ('E09 · C-FALDA CON TAJALÍ', 'Calza-falda con tajalí', 'Varios', 'Falda', 'Femenino', 'Calza-Falda', [('Falda delantero', 1), ('Falda trasero', 1), ('Calza pierna', 2), ('Pretina', 1)]),
]
PZ_NUEVAS = {n[0]: n[6] for n in NUEVOS}
_sm = [x for x in rec6 if str(x[0]).startswith('SM-')]
rec6 = [x for x in rec6 if not str(x[0]).startswith('SM-')] + \
       [(c, d, dep, pr, ca, None, None, None, None, None, None, None, cf, None, None, 'Por confirmar', 'Agregado desde el catálogo vigente: piezas sugeridas, revisar en la marcada.')
        for c, d, dep, pr, ca, cf, _ in NUEVOS] + _sm
p6 = v6['Parametros']

F0 = 5                      # primera fila de datos (encabezado en la fila 4)
NT, NC, NA = 120, 80, 60    # filas disponibles: telas, confección, avíos
NM, ND = 150, 1000          # moldes, piezas

# ================================================================ PRECIOS
wp = wb.create_sheet('Precios')
titulo(wp, 'PRECIOS — telas, confección, avíos e impresión', 'Azul = lo escribes tú. Gris = se calcula solo. Cada tela se anota como se compra (kg o metro) con el precio de la factura; con los kilos por metro se convierte a $ por kg y $ por metro, netos.')
# Telas A:L — se anota como se compra (kg o metro) y se convierte a $ por kg y $ por metro, netos
enc(wp, 4, 1, ['TELA', 'Tipo', 'Proveedor', 'Se compra en\n(kg / metro)', 'PRECIO ACTUAL\n(factura)', '¿Incluye\nIVA?', 'Kilos por metro\n(1 m lineal)', 'Ancho útil\n(m)',
               '$ POR KG\nneto', '$ POR METRO\nneto', '$ POR M²\nneto', 'Nota'], [30, 10, 20, 11, 12, 9, 12, 9, 11, 12, 10, 40])
for i in range(NT):
    r = F0 + i
    if i < len(TELAS):
        t = TELAS[i]
        act = t.get('act') if t.get('act') is not None else t.get('reg')
        for col, v in ((1, t['n']), (2, t.get('tipo')), (3, t.get('prov')), (4, t.get('u')), (5, act),
                       (6, t.get('iva')), (7, t.get('kgm')), (8, t.get('an'))):
            wp.cell(r, col, v)
        nota = ' '.join(x for x in (t.get('nota'), ('REVISAR: ' + t['duda']) if t.get('duda') else None,
                                    None if t.get('act') is not None else 'Precio del sistema anterior SIN CONFIRMAR (¿por kg o por metro?).') if x)
        wp.cell(r, 12, nota or None)
    for c in (1, 2, 3, 4, 5, 6, 7, 8, 12): wp.cell(r, c).fill = AZUL
    wp.cell(r, 5).number_format = MONEY; wp.cell(r, 7).number_format = '0.000'; wp.cell(r, 8).number_format = '0.00'
    neto = f'E{r}/IF(LOWER(LEFT(F{r},1))="s",1.19,1)'
    wp.cell(r, 9, f'=IF(NOT(ISNUMBER(E{r})),"",IF(LOWER(D{r})="metro",IF(N(G{r})>0,{neto}/G{r},""),{neto}))').number_format = MONEY
    wp.cell(r, 10, f'=IF(NOT(ISNUMBER(E{r})),"",IF(LOWER(D{r})="metro",{neto},IF(N(G{r})>0,{neto}*G{r},"")))').number_format = MONEY
    wp.cell(r, 11, f'=IF(AND(ISNUMBER(J{r}),N(H{r})>0),J{r}/H{r},"")').number_format = MONEY
    for c in (9, 10, 11): wp.cell(r, c).fill = GRIS
    wp.cell(r, 10).font = B
    wp.cell(r, 12).alignment = Alignment(wrap_text=False)
lista(wp, '"Nacional,Pro,Importada"', f'B{F0}:B{F0+NT-1}')
lista(wp, '"kg,metro"', f'D{F0}:D{F0+NT-1}')
lista(wp, '"Sí,No"', f'F{F0}:F{F0+NT-1}')
TN, TPM, TAN, TKGM = (f'Precios!${c}${F0}:${c}${F0+NT-1}' for c in 'AJHG')
# Confección N:O
enc(wp, 4, 14, ['CONFECCIÓN (tipo)', '$ por prenda\n(taller)'], [34, 12])
for i in range(NC):
    r = F0 + i
    if i < len(conf6): wp.cell(r, 14, conf6[i][0]); wp.cell(r, 15, conf6[i][1])
    wp.cell(r, 14).fill = wp.cell(r, 15).fill = AZUL; wp.cell(r, 15).number_format = MONEY
CN, CP = f'Precios!$N${F0}:$N${F0+NC-1}', f'Precios!$O${F0}:$O${F0+NC-1}'
# Avíos Q:R
enc(wp, 4, 17, ['AVÍO', '$ por unidad'], [22, 12])
for i in range(NA):
    r = F0 + i
    if i < len(avios6): wp.cell(r, 17, avios6[i][0]); wp.cell(r, 18, avios6[i][3])
    wp.cell(r, 17).fill = wp.cell(r, 18).fill = AZUL; wp.cell(r, 18).number_format = MONEY
AN, AP = f'Precios!$Q${F0}:$Q${F0+NA-1}', f'Precios!$R${F0}:$R${F0+NA-1}'
# Impresión y parámetros T:V
enc(wp, 4, 20, ['IMPRESIÓN Y PARÁMETROS', 'Valor', 'Nota'], [40, 13, 46])
PAR = [('Papel sublimación: precio del rollo', p6['B21'].value, MONEY, 'POR CONFIRMAR: ¿$50.420 es el rollo?', True),
       ('Papel sublimación: metros del rollo', p6['B22'].value, None, 'PENDIENTE', True),
       ('Tinta: precio por litro', p6['B23'].value, MONEY, 'POR CONFIRMAR', True),
       ('Tinta: ml por metro impreso', p6['B24'].value, None, 'PENDIENTE: litros del mes ÷ metros impresos × 1000', True),
       ('Papel protección (kraft): $ por metro', None, MONEY, 'Opcional', True),
       ('COSTO DE IMPRESIÓN POR METRO', '=IFERROR(U5/U6,0)+N(U8)*N(U7)/1000+N(U9)', MONEY, 'Papel + tinta + protección', False),
       ('', None, None, '', None),
       ('Merma general sobre el despiece', 0, PCT, 'Tela que se pierde (cabezales, fallas). Se usa si el molde no tiene ajuste propio.', True),
       ('', None, None, '', None),
       ('Gastos generales (% del precio)', 0.434, PCT, 'Del sistema de ventas (jul–sep 2026). Solo para el precio mínimo.', True),
       ('Meta de ganancia (% del precio)', 0.15, PCT, 'POR CONFIRMAR', True),
       ('Multiplicador: precio mínimo = costo ×', '=IFERROR(1/(1-U14-U15),"")', '0.00', 'Costo ÷ (1 − gastos generales − meta)', False)]
for i, (a, v, fmt, nota, inp) in enumerate(PAR):
    r = F0 + i
    if not a: continue
    wp.cell(r, 20, a).font = Font(bold=not inp)
    c = wp.cell(r, 21, v); c.fill = AZUL if inp else AMA
    if fmt: c.number_format = fmt
    if not inp: c.font = B
    wp.cell(r, 22, nota).font = Font(color='595959', size=9)
IMP, MERMA, GG, META, MULT = 'Precios!$U$10', 'Precios!$U$12', 'Precios!$U$14', 'Precios!$U$15', 'Precios!$U$16'
wp.freeze_panes = 'A5'

# ================================================================ DESPIECE
def piezas(code, desc, prenda):
    d = (code + ' ' + desc).upper(); last = code.split()[-1]
    pro, rev = 'PRO' in d, (last.endswith('RE') or 'REVERSIBLE' in d)
    rag = 'RAGL' in d or bool(re.search(r'02(PRO)?$', last))
    if prenda == 'Short':
        p = [('Delantero', 2), ('Trasero', 2), ('Pretina', 1)]
        if rev: p = [(n + ' cara A', q) for n, q in p] + [(n + ' cara B', q) for n, q in p]
        return p + ([('Paneles laterales', 2)] if pro else [])
    if prenda == 'Calza': return [('Pierna', 2), ('Pretina', 1)]
    if 'BK' in code or 'RUNNING' in d:
        p = [('Delantero', 1), ('Espalda', 1), ('Ribete cuello y sisas', 1)]
        if rev: p = [('Delantero cara A', 1), ('Espalda cara A', 1), ('Delantero cara B', 1), ('Espalda cara B', 1), ('Ribete cuello y sisas', 1)]
        return p + ([('Paneles laterales', 2)] if pro else [])
    p = [('Delantero', 1), ('Espalda', 1), ('Manga raglán' if rag else 'Manga', 2), ('Tapeta' if 'CAMISERO' in d else 'Tira de cuello', 1)]
    return p + ([('Paneles laterales', 2)] if pro else [])

wd = wb.create_sheet('Despiece')
titulo(wd, 'DESPIECE — piezas de cada molde', 'Metros por prenda de cada pieza (de la marcada del plotter). Los metros NO cambian con la tela. "Tela por defecto" es la que usa la Ficha si no eliges otra.')
enc(wd, 4, 1, ['Molde', 'Pieza', 'Cant.', 'METROS\npor prenda\n(en el ancho\nde referencia)', 'Tela por defecto', '¿Se\nsublima?', '$ tela + impresión\n(con tela por defecto)', 'clave'], [24, 24, 7, 11, 30, 8, 16, 6])
fila = F0
for x in rec6:
    code, t1, t2 = x[0], x[5], x[8]
    for n, q in (PZ_NUEVAS.get(code) or piezas(code, x[1] or '', x[3] or '')):
        wd.cell(fila, 1, code); wd.cell(fila, 2, n); wd.cell(fila, 3, q)
        wd.cell(fila, 5, t2 if (n == 'Paneles laterales' and t2) else t1)
        wd.cell(fila, 6, 'No' if code.startswith(('SM-02', 'SM-04')) else 'Sí')
        fila += 1
for r in range(F0, F0 + ND):
    for c in range(1, 7): wd.cell(r, c).fill = AZUL
    wd.cell(r, 4).number_format = M3
    aj = f'(1+IFERROR(IF(ISNUMBER(INDEX(Moldes!$K${F0}:$K${F0+NM-1},MATCH(A{r},Moldes!$A${F0}:$A${F0+NM-1},0))),INDEX(Moldes!$K${F0}:$K${F0+NM-1},MATCH(A{r},Moldes!$A${F0}:$A${F0+NM-1},0)),{MERMA}),{MERMA}))'
    ref = f'INDEX(Moldes!$R${F0}:$R${F0+NM-1},MATCH(A{r},Moldes!$A${F0}:$A${F0+NM-1},0))'
    fac = f'IFERROR(IF(AND(N({ref})>0,N(INDEX({TAN},MATCH(E{r},{TN},0)))>0),{ref}/INDEX({TAN},MATCH(E{r},{TN},0)),1),1)'
    c = wd.cell(r, 7, f'=IF(OR(A{r}="",N(D{r})=0),"",IFERROR(D{r}*{aj}*(INDEX({TPM},MATCH(E{r},{TN},0))*{fac}+IF(F{r}="Sí",{IMP},0)),"falta $ tela"))')
    c.number_format = MONEY; c.fill = GRIS
    c = wd.cell(r, 8, f'=IF(A{r}="","",A{r}&"|"&COUNTIF($A${F0}:A{r},A{r}))'); c.font = Font(color='BBBBBB', size=8)
lista(wd, f'=Moldes!$A${F0}:$A${F0+NM-1}', f'A{F0}:A{F0+ND-1}')
lista(wd, f'={TN}', f'E{F0}:E{F0+ND-1}')
lista(wd, '"Sí,No"', f'F{F0}:F{F0+ND-1}')
wd.freeze_panes = 'C5'; wd.auto_filter.ref = f'A4:G{F0+ND-1}'
DA, DB_, DC, DD, DE, DF, DG, DK = (f'Despiece!${c}${F0}:${c}${F0+ND-1}' for c in 'ABCDEFGH')

# ================================================================ MOLDES
wm = wb.create_sheet('Moldes')
titulo(wm, 'MOLDES — catálogo con su costo de producción (telas por defecto)', 'Una fila por molde. Lo que va por prenda y no depende de la tela: confección y avíos. Para probar otras telas, usa la Ficha.')
enc(wm, 4, 1, ['Molde', 'Descripción', 'Prenda', 'Calce', 'Tipo de confección', 'Avío 1', 'Cant.', 'Avío 2', 'Cant.', 'Piezas\nen despiece',
               'Ajuste por consumo\nreal (%) — vacío =\nmerma general', '$ Tela +\nimpresión', '$ Confección', '$ Avíos', 'COSTO DE\nPRODUCCIÓN', 'Qué falta', 'Notas', 'Ancho de la tela\ncon que se midió\nel despiece (m)'],
    [24, 30, 9, 15, 30, 13, 6, 13, 6, 9, 16, 12, 12, 10, 13, 40, 60])
NOTAS_MED = {'04 · FUTINHOMU01': 'Corte OT 2369: 0,0375 kg Cire + 0,024 kg Mesh por prenda.', '13 · FUTMUJU01PRO': 'Corte OT 2257: 0,1225 kg Montesimone por prenda.',
             '14 · FUTHOMU01PRO': 'Corte OT 2547: 0,12 kg Montesimone + 0,0775 kg Mesh por prenda.', '30 · BKHOMU03': 'Corte OT 2425: 0,2029 kg Quickdry por prenda.',
             '46 · BK-SH': 'Corte OT 2425: 0,1529 kg Quickdry por prenda.', 'SM-01 · Short Fútbol': 'Cortes OT 2547 y 2588: 0,175 y 0,138 kg Cire por prenda.',
             'SM-02 · Short Rugby': 'Corte OT 205: 0,1842 kg Gabardina por prenda.', 'SM-03 · Camiseta Running s/m': 'Corte OT 2223: 0,082 kg 5K por prenda.',
             'SM-04 · Calza Vóley': 'Corte OT 2642: 0,074 kg Lycra por prenda.'}
for i in range(NM):
    r = F0 + i
    if i < len(rec6):
        x = rec6[i]
        wm.cell(r, 1, x[0]); wm.cell(r, 2, x[1]); wm.cell(r, 3, x[3]); wm.cell(r, 4, x[4]); wm.cell(r, 5, x[12])
        wm.cell(r, 6, x[13]); wm.cell(r, 7, x[14]); wm.cell(r, 17, NOTAS_MED.get(x[0]))
    for c in (1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 17, 18): wm.cell(r, c).fill = AZUL
    wm.cell(r, 18).number_format = '0.00'
    wm.cell(r, 11).number_format = PCT
    wm[f'J{r}'] = f'=IF(A{r}="","",COUNTIF({DA},A{r}))'
    wm[f'L{r}'] = f'=IF(A{r}="","",IF(COUNTIFS({DA},A{r},{DD},">0")<J{r},"",IF(COUNTIFS({DA},A{r},{DG},"falta*")>0,"",SUMIFS({DG},{DA},A{r}))))'
    wm[f'M{r}'] = f'=IF(OR(A{r}="",E{r}=""),"",IFERROR(INDEX({CP},MATCH(E{r},{CN},0)),""))'
    wm[f'N{r}'] = f'=IF(A{r}="","",IFERROR(N(G{r})*IF(F{r}="",0,INDEX({AP},MATCH(F{r},{AN},0)))+N(I{r})*IF(H{r}="",0,INDEX({AP},MATCH(H{r},{AN},0))),""))'
    wm[f'O{r}'] = f'=IF(A{r}="","",IFERROR(L{r}+M{r}+N{r},""))'
    wm[f'P{r}'] = (f'=IF(A{r}="","",IF(ISNUMBER(O{r}),"Lista"&IF(N({IMP})=0," (impresión en $0)",""),'
                   f'IF(J{r}=0,"Agregar piezas en Despiece. ",IF(COUNTIFS({DA},A{r},{DD},">0")<J{r},"Metros de "&(J{r}-COUNTIFS({DA},A{r},{DD},">0"))&" pieza(s). ",""))'
                   f'&IF(COUNTIFS({DA},A{r},{DG},"falta*")>0,"Precio o peso de 1 m de alguna tela. ","")&IF(M{r}="","Tipo/precio de confección. ","")))')
    for c, fmt in (('L', MONEY), ('M', MONEY), ('N', MONEY), ('O', MONEY)): wm[f'{c}{r}'].number_format = fmt
    for c in 'JLMNP': wm[f'{c}{r}'].fill = GRIS
    wm[f'O{r}'].fill = AMA; wm[f'O{r}'].font = B; wm[f'P{r}'].font = ROJO
lista(wm, f'={CN}', f'E{F0}:E{F0+NM-1}')
for c in 'FH': lista(wm, f'={AN}', f'{c}{F0}:{c}{F0+NM-1}')
wm['A3'] = f'="Moldes con costo completo: "&COUNTIF(P{F0}:P{F0+NM-1},"Lista*")&" de "&COUNTA(A{F0}:A{F0+NM-1})'
wm['A3'].font = Font(bold=True, color='C00000')
wm.freeze_panes = 'B5'
MA = f'Moldes!$A${F0}:$A${F0+NM-1}'
def MX(col): return f'INDEX(Moldes!${col}${F0}:${col}${F0+NM-1},MATCH($C$4,{MA},0))'

# ================================================================ FICHA
wf = wb.active; wf.title = 'Ficha'
wf['A1'] = 'FICHA DE COSTO DE PRODUCCIÓN'; wf['A1'].font = Font(bold=True, size=16, color='1F3864')
wf['A2'] = 'Elige el molde y, si quieres, otra tela para cada pieza. Vacío = tela por defecto. Todo neto, sin IVA y sin gastos generales.'; wf['A2'].font = Font(italic=True, color='595959')
for col, w in zip('ABCDEFGHIJ', [4, 24, 7, 10, 30, 30, 12, 12, 12, 13]): wf.column_dimensions[col].width = w
wf.column_dimensions['K'].width = 40
wf['B4'] = 'MOLDE ▶'; wf['B4'].font = Font(bold=True, size=13, color='1F3864')
wf.merge_cells('C4:E4'); c = wf['C4']; c.value = '14 · FUTHOMU01PRO'; c.fill = AZUL; c.font = Font(bold=True, size=13)
lista(wf, f'={MA}', 'C4')
wf['F4'] = f'=IFERROR({MX("B")}&"  ·  "&{MX("C")}&" "&{MX("D")},"← elige un molde")'; wf['F4'].font = Font(italic=True, color='595959')
wf['B5'] = f'=IFERROR(IF(N({MX("R")})>0,"Despiece medido en tela de "&TEXT({MX("R")},"0.00")&" m de ancho: los metros se ajustan al ancho de cada tela.  ","Sin ancho de referencia (Moldes): los metros no se ajustan por ancho.  ")&"Consumo: despiece + "&TEXT(IF(ISNUMBER({MX("K")}),{MX("K")},{MERMA}),"0%")&IF(ISNUMBER({MX("K")})," (ajuste por consumo real de este molde)"," (merma general)"),"")'
wf['B5'].font = Font(italic=True, color='595959', size=9)
wf['B3'] = 'Al cambiar de molde, borra las telas elegidas (columna azul) para volver a las de siempre.'; wf['B3'].font = Font(italic=True, color='C00000', size=9)
H0, NP = 7, 16
for i, t in enumerate(['#', 'Pieza', 'Cant.', 'Metros\npor prenda', 'TELA (elige)\nvacío = por defecto', 'Tela usada', '$ por metro', '$ Tela', '$ Impresión', 'TOTAL PIEZA', 'Aviso'], 1):
    c = wf.cell(H0, i, t); c.font = H; c.fill = HF; c.alignment = WC
wf.row_dimensions[H0].height = 40
AJ = f'(1+IFERROR(IF(ISNUMBER({MX("K")}),{MX("K")},{MERMA}),{MERMA}))'
for i in range(NP):
    r = H0 + 1 + i; k = f'$C$4&"|"&{i+1}'; idx = f'MATCH({k},{DK},0)'
    wf[f'A{r}'] = f'=IF(B{r}="","",{i+1})'
    wf[f'B{r}'] = f'=IFERROR(INDEX({DB_},{idx}),"")'
    wf[f'C{r}'] = f'=IF(B{r}="","",INDEX({DC},{idx}))'
    fac = f'IFERROR(IF(AND(N({MX("R")})>0,N(INDEX({TAN},MATCH(F{r},{TN},0)))>0),{MX("R")}/INDEX({TAN},MATCH(F{r},{TN},0)),1),1)'
    wf[f'D{r}'] = f'=IF(B{r}="","",IF(N(INDEX({DD},{idx}))>0,INDEX({DD},{idx})*{AJ}*{fac},""))'
    wf[f'E{r}'].fill = AZUL
    wf[f'F{r}'] = f'=IF(B{r}="","",IF(E{r}<>"",E{r},INDEX({DE},{idx})&""))'
    wf[f'G{r}'] = f'=IF(OR(B{r}="",F{r}=""),"",IFERROR(INDEX({TPM},MATCH(F{r},{TN},0))*1,""))'
    wf[f'H{r}'] = f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(G{r})),D{r}*G{r},"")'
    wf[f'I{r}'] = f'=IF(OR(B{r}="",NOT(ISNUMBER(D{r}))),"",IF(INDEX({DF},{idx})="Sí",INDEX({DD},{idx})*{AJ}*{IMP},0))'
    wf[f'J{r}'] = f'=IF(AND(ISNUMBER(H{r}),ISNUMBER(I{r})),H{r}+I{r},"")'
    wf[f'K{r}'] = (f'=IF(B{r}="","",IF(ISNUMBER(J{r}),IF(INDEX({DF},{idx})="Sí","Sublimada","")&IF(ABS({fac}-1)>0.001,"  · ajustado al ancho "&TEXT(INDEX({TAN},MATCH(F{r},{TN},0)),"0.00")&" m",""),'
                   f'IF(NOT(ISNUMBER(D{r})),"Faltan metros (Despiece). ","")&IF(F{r}="","Elige tela. ",IF(NOT(ISNUMBER(G{r})),"Falta precio o peso de 1 m (Precios). ",""))))')
    for c, fmt in (('D', M3), ('G', MONEY), ('H', MONEY), ('I', MONEY), ('J', MONEY)): wf[f'{c}{r}'].number_format = fmt
    for c in 'ABCDFGHI': wf[f'{c}{r}'].fill = GRIS
    wf[f'J{r}'].fill = AMA; wf[f'J{r}'].font = B; wf[f'K{r}'].font = ROJO
    lista(wf, f'={TN}', f'E{r}')
last = H0 + NP; P1, PN = H0 + 1, last
r = last + 2
wf.cell(r, 2, 'RESUMEN POR PRENDA').font = Font(bold=True, size=12, color='1F3864')
for c in range(2, 11): wf.cell(r, c).border = LINEA
completo = f'SUMPRODUCT((B{P1}:B{PN}<>"")*NOT(ISNUMBER(J{P1}:J{PN})))=0'
res = [('Tela', f'=IF({completo},SUM(H{P1}:H{PN}),"Faltan datos ↑")', 'Suma de las piezas'),
       ('Papel y tinta de sublimación', f'=IF({completo},SUM(I{P1}:I{PN}),"Faltan datos ↑")', f'=IF(N({IMP})=0,"OJO: costo de impresión en $0 (hoja Precios)",TEXT(SUMIF(I{P1}:I{PN},">0",D{P1}:D{PN}),"0.00")&" m impresos × "&TEXT({IMP},"$#,##0"))'),
       ('Confección (taller)', f'=IFERROR({MX("M")}*1,"Falta tipo de confección (Moldes)")', f'=IFERROR({MX("E")}&"","")'),
       ('Avíos', f'=IFERROR({MX("N")}*1,0)', f'=IFERROR(IF({MX("F")}&""="","Sin avíos",{MX("F")}&" × "&{MX("G")}&IF({MX("H")}&""="",""," + "&{MX("H")}&" × "&{MX("I")})),"")'),
       ('Extras (bordado, DTF, escudo…) — escribe', None, 'Por prenda')]
for j, (lab, f, nota) in enumerate(res):
    rr = r + 1 + j
    wf.cell(rr, 2, lab); c = wf.cell(rr, 10, f); c.number_format = MONEY; c.fill = AZUL if f is None else GRIS
    wf.cell(rr, 11, nota).font = Font(color='595959', size=9)
rc = r + 1 + len(res)
wf.cell(rc, 2, 'COSTO DE PRODUCCIÓN').font = Font(bold=True, size=13)
c = wf.cell(rc, 10, f'=IFERROR(J{r+1}+J{r+2}+J{r+3}+J{r+4}+N(J{r+5}),"Completa los datos ↑")'); c.number_format = MONEY; c.fill = AMA; c.font = Font(bold=True, size=13)
for cc in range(2, 11): wf.cell(rc, cc).border = LINEA
rp = rc + 2
wf.cell(rp, 2, f'="Precio mínimo de venta (neto) = costo × "&TEXT({MULT},"0.00")&"   (gastos generales "&TEXT({GG},"0.0%")&" + ganancia "&TEXT({META},"0%")&")"')
c = wf.cell(rp, 10, f'=IF(ISNUMBER(J{rc}),J{rc}*{MULT},"")'); c.number_format = MONEY; c.fill = AMA; c.font = B
wf.cell(rp + 1, 2, 'Tu precio de venta actual, neto — escribe')
c = wf.cell(rp + 1, 10); c.number_format = MONEY; c.fill = AZUL
wf.cell(rp + 2, 2, 'Ganancia por prenda (después de costo y gastos generales)')
c = wf.cell(rp + 2, 10, f'=IF(AND(ISNUMBER(J{rc}),N(J{rp+1})>0),J{rp+1}-J{rc}-J{rp+1}*{GG},"")'); c.number_format = MONEY; c.fill = GRIS; c.font = B
wf.cell(rp + 2, 11, f'=IF(ISNUMBER(J{rp+2}),TEXT(J{rp+2}/J{rp+1},"0.0%")&" del precio"&IF(J{rp+2}<0,"  ← PIERDE PLATA",IF(J{rp+2}/J{rp+1}<{META},"  ← bajo la meta","")),"")').font = B
# calibrar con un corte real
rk = rp + 4
wf.cell(rk, 2, 'AJUSTAR CON UN CORTE REAL (opcional)').font = Font(bold=True, size=12, color='1F3864')
for cc in range(2, 11): wf.cell(rk, cc).border = LINEA
wf.cell(rk + 1, 2, 'Después de cortar este molde: qué tela, cuántas prendas y cuántos kg se usaron. Te dice el % de ajuste para anotar en Moldes.').font = Font(italic=True, color='595959', size=9)
cal = [('Tela cortada', None, None), ('Prendas cortadas', None, '0'), ('Kg usados', None, '0.00'),
       ('Metros reales por prenda', f'=IFERROR(J{rk+4}/INDEX({TKGM},MATCH(J{rk+2},{TN},0))/J{rk+3},"")', M3),
       ('Metros del despiece en esa tela (sin merma)', f'=IF(J{rk+2}="","",SUMIF(F{P1}:F{PN},J{rk+2},D{P1}:D{PN})/{AJ})', M3),
       ('AJUSTE A ANOTAR EN MOLDES (columna K)', f'=IF(AND(ISNUMBER(J{rk+5}),N(J{rk+6})>0),J{rk+5}/J{rk+6}-1,"")', PCT)]
for j, (lab, f, fmt) in enumerate(cal):
    rr = rk + 2 + j
    wf.cell(rr, 2, lab).font = Font(bold=(j == 5))
    c = wf.cell(rr, 10, f); c.fill = AZUL if f is None else (AMA if j == 5 else GRIS)
    if fmt: c.number_format = fmt
lista(wf, f'={TN}', f'J{rk+2}')
wf.cell(rk + 8, 2, 'Si da error: falta el peso de 1 metro de esa tela, o los metros de las piezas en el Despiece.').font = Font(italic=True, color='595959', size=9)
wf.cell(rk + 10, 2, 'No van aquí (están en los gastos generales del sistema de ventas): sueldos de corte, ploteo, sublimación, estampado y diseño, arriendo, luz, ventas y administración.').font = Font(italic=True, color='595959', size=9)
wf.freeze_panes = 'A8'

wb._sheets = [wb['Ficha'], wb['Moldes'], wb['Despiece'], wb['Precios']]
wb.save(DST)
print('ok', len(rec6), 'moldes', fila - F0, 'piezas', len(telas6), 'telas')
