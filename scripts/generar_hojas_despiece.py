# Hojas de despiece (PDF rellenable, una página por molde) para llevar al plotter y al corte.
# Los nombres de campo son los mismos que lee importar_formulario_receta.py.
import sys, re
from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, black, white

SRC, DST = sys.argv[1], sys.argv[2]
PEDIDOS = sys.argv[3].split(',')          # números de molde, ej: 01,02,29,30,46,61
wb = load_workbook(SRC)
M, D, P = wb['Moldes'], wb['Despiece'], wb['Precios']
moldes = {str(M.cell(r, 1).value).split(' ')[0]: [M.cell(r, c).value for c in range(1, 19)] for r in range(5, 155) if M.cell(r, 1).value}
piezas = {}
for r in range(5, 1005):
    if D.cell(r, 1).value: piezas.setdefault(D.cell(r, 1).value, []).append([D.cell(r, c).value for c in range(1, 7)])
telas = [P.cell(r, 1).value for r in range(5, 125) if P.cell(r, 1).value]
AZUL = HexColor('#1F3864'); GRIS = HexColor('#595959'); FONDO = HexColor('#EAF1FB'); LINEA = HexColor('#BFBFBF')
W, H = letter; X0, X1 = 36, W - 36
c = canvas.Canvas(DST, pagesize=letter); c.setTitle('PMK - Hojas de despiece'); form = c.acroForm
def san(s): return re.sub(r'[^A-Za-z0-9]+', '_', str(s or '')).strip('_')[:40]
def num(v):
    if v in (None, ''): return ''
    return (f'{v:.3f}'.rstrip('0').rstrip('.')).replace('.', ',') if isinstance(v, float) else str(v)
def campo(n, x, y, w, v='', h=14, fs=8.5):
    form.textfield(name=n, tooltip=n, x=x, y=y - 4, width=w, height=h, value=str(v or ''), fontName='Helvetica', fontSize=fs,
                   borderWidth=0.5, borderColor=LINEA, fillColor=FONDO, textColor=black, forceBorder=True)
def txt(x, y, s, size=8.5, bold=False, color=black):
    c.setFont('Helvetica-Bold' if bold else 'Helvetica', size); c.setFillColor(color); c.drawString(x, y, s); c.setFillColor(black)
def sec(y, t):
    c.setFillColor(HexColor('#D9D9D9')); c.rect(X0, y - 4, X1 - X0, 15, fill=1, stroke=0); txt(X0 + 4, y, t, 9, True); return y - 22
ncorte = 0
for num_m in PEDIDOS:
    v = moldes.get(num_m)
    code = v[0] if v else None
    key = san(code) if code else num_m
    c.setFillColor(AZUL); c.rect(0, H - 54, W, 54, fill=1, stroke=0)
    txt(X0, H - 30, f'HOJA DE DESPIECE  -  MOLDE {num_m}', 15, True, white)
    txt(X0, H - 45, 'Completar en el plotter (marcada) y en el corte. Devolver el PDF o una foto para actualizar la receta.', 8.5, False, white)
    y = H - 76
    if v:
        txt(X0, y, f'{code}', 13, True, AZUL); txt(X0 + 170, y + 1, f'{v[1] or ""}  -  {v[2] or ""} {v[3] or ""}', 9.5, False, GRIS); y -= 22
    else:
        txt(X0, y, 'Molde nuevo (no esta en la receta todavia):', 9.5, True)
        txt(X0 + 205, y, 'Codigo:', 8.5); campo(f'MOLDENUEVO_{num_m}__codigo', X0 + 240, y, 140, f'{num_m} · ')
        txt(X0 + 390, y, 'Prenda:', 8.5); campo(f'MOLDENUEVO_{num_m}__prenda', X0 + 425, y, 115); y -= 20
        txt(X0, y, 'Descripcion:', 8.5); campo(f'MOLDENUEVO_{num_m}__descripcion', X0 + 60, y, 250)
        txt(X0 + 320, y, 'Calce:', 8.5); campo(f'MOLDENUEVO_{num_m}__calce', X0 + 350, y, 190); y -= 24
    mk = f'MOLDE__{key}'
    txt(X0, y, 'Responsable:', 8.5); campo(f'{mk}__responsable', X0 + 62, y, 150); txt(X0 + 225, y, 'Fecha:', 8.5); campo(f'{mk}__fecha', X0 + 255, y, 80)
    txt(X0 + 350, y, 'Talla base:', 8.5); campo(f'{mk}__talla', X0 + 400, y, 50, 'M'); y -= 26

    y = sec(y, '1. MARCADAS EN EL PLOTTER (talla base) - una por cada tela del molde')
    lst = (piezas.get(code) or []) if code else []
    tl = []
    for p in lst:
        if p[4] and p[4] not in tl: tl.append(p[4])
    for x, t in ((X0 + 2, 'N°'), (X0 + 22, 'Tela de la marcada'), (X0 + 222, 'Ancho util (m)'), (X0 + 292, 'Metros marcada'), (X0 + 362, 'Prendas'), (X0 + 412, 'Aprovechamiento (%)')): txt(x, y, t, 7.5, True)
    y -= 15
    for n in (1, 2, 3):
        txt(X0 + 6, y, str(n), 9, True)
        campo(f'{mk}__m{n}_tela', X0 + 20, y, 196, tl[n - 1] if len(tl) >= n else ''); campo(f'{mk}__m{n}_ancho', X0 + 220, y, 66)
        campo(f'{mk}__m{n}_metros', X0 + 290, y, 66); campo(f'{mk}__m{n}_prendas', X0 + 360, y, 46); campo(f'{mk}__m{n}_aprov', X0 + 410, y, 70)
        y -= 17
    txt(X0, y, 'Cada tela se marca por separado y tiene su propio ancho util. Metros por prenda = metros de la marcada / prendas.', 7.5, False, GRIS); y -= 20

    y = sec(y, '2. PIEZAS - metros por prenda de cada fila (ya sumando la cantidad), medidos en su marcada')
    for x, t in ((X0 + 2, 'Pieza'), (X0 + 152, 'Cant.'), (X0 + 186, 'Marcada N°'), (X0 + 236, 'Metros por prenda'), (X0 + 316, 'Tela por defecto'), (X0 + 492, 'Sublima')): txt(x, y, t, 7.5, True)
    y -= 15
    for j, p in enumerate(lst + [[None, '', '', '', '', '']] * (10 - len(lst) if len(lst) < 10 else 2)):
        k = f'PZA__{key}__{j+1:02d}'
        mnum = (tl.index(p[4]) + 1) if (p[4] and p[4] in tl) else (1 if p[1] else '')
        campo(f'{k}__pieza', X0, y, 148, p[1]); campo(f'{k}__cant', X0 + 150, y, 32, p[2]); campo(f'{k}__marcada', X0 + 186, y, 44, mnum)
        campo(f'{k}__metros', X0 + 234, y, 76, num(p[3])); campo(f'{k}__tela', X0 + 314, y, 172, p[4]); campo(f'{k}__sublima', X0 + 490, y, 50, p[5] or ('Sí' if not code else ''))
        y -= 17
    txt(X0, y, 'Suma de metros por marcada (control): debe parecerse a metros de la marcada / prendas.', 7.5, False, GRIS); y -= 24

    y = sec(y, '3. CORTE REAL (opcional, para ajustar la merma) - pesando la tela antes y despues')
    for x, t in ((X0 + 2, 'Fecha'), (X0 + 62, 'OT'), (X0 + 120, 'Tela cortada'), (X0 + 330, 'Prendas'), (X0 + 385, 'Kg antes'), (X0 + 440, 'Kg despues'), (X0 + 495, 'Kg usados')): txt(x, y, t, 7.5, True)
    y -= 15
    for i in range(3):
        ncorte += 1; k = f'CORTE_{ncorte:03d}'
        campo(f'{k}__fecha', X0, y, 56); campo(f'{k}__ot', X0 + 60, y, 54); campo(f'{k}__tela', X0 + 118, y, 205); campo(f'{k}__prendas', X0 + 328, y, 50)
        campo(f'{k}__kg_antes', X0 + 383, y, 50); campo(f'{k}__kg_despues', X0 + 438, y, 50); campo(f'{k}__kg', X0 + 493, y, 47)
        form.textfield(name=f'{k}__molde', tooltip='molde', x=-100, y=-100, width=1, height=1, value=(code or num_m))   # oculto: molde de esta hoja
        y -= 17
    txt(X0, y, 'Kg usados = kg antes - kg despues (lo que vuelve al rollo como tela util). Los retazos cuentan como usados.', 7.5, False, GRIS); y -= 20
    txt(X0, y, 'Observaciones:', 8.5); campo(f'{mk}__obs', X0 + 70, y - 14, X1 - X0 - 70, '', h=28)
    c.setFont('Helvetica', 7); c.setFillColor(GRIS); c.drawString(X0, 20, 'Playmaker - Hoja de despiece. Telas: usar el nombre exacto del listado de telas.'); c.setFillColor(black)
    c.showPage()
c.save(); print('ok', len(PEDIDOS), 'hojas')
