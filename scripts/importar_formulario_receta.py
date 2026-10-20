"""Pasa los datos de PMK_Formulario_Receta.pdf (Precios: telas A-L con D=se compra en, E=precio, F=IVA, G=kg/m, H=ancho; confección N:O; avíos Q:R; parámetros U. Moldes R=ancho de referencia) (completado) a PMK_Receta_Produccion.xlsx.

Uso: python3 importar_formulario_receta.py formulario.pdf PMK_Receta_Produccion.xlsx salida.xlsx
Imprime cada cambio hecho. Solo escribe en celdas de entrada (azules); las fórmulas no se tocan.
"""
import re, sys, collections
from pypdf import PdfReader
from openpyxl import load_workbook

PDF, XLSX, OUT = sys.argv[1:4]
f = {k: (v.get('/V') or '').strip() for k, v in (PdfReader(PDF).get_fields() or {}).items() if isinstance(v.get('/V'), str)}
wb = load_workbook(XLSX)
P, M, D = wb['Precios'], wb['Moldes'], wb['Despiece']
cambios = []

def num(s):
    s = (s or '').replace('$', '').replace(' ', '').replace('%', '')
    if not s: return None
    if ',' in s: s = s.replace('.', '').replace(',', '.')
    elif s.count('.') == 1 and len(s.split('.')[1]) == 3 and s.split('.')[0].lstrip('-') not in ('', '0'): s = s.replace('.', '')  # 5.600 = cinco mil seiscientos; 0.150 = decimal
    else: s = s.replace('.', '') if s.count('.') > 1 else s
    try: return float(s) if '.' in s else int(s)
    except ValueError: return None
def san(s): return re.sub(r'[^A-Za-z0-9]+', '_', str(s or '')).strip('_')[:40]
def poner(ws, r, c, v, que):
    if v is None or v == '': return
    old = ws.cell(r, c).value
    if old != v and not (isinstance(old, (int, float)) and isinstance(v, (int, float)) and abs(old - v) < 1e-9):
        ws.cell(r, c).value = v; cambios.append(f'{ws.title} fila {r}: {que}: {old!r} -> {v!r}')
def fila_de(ws, col, nombre, r0, r1):
    n = (nombre or '').strip().upper()
    for r in range(r0, r1):
        if str(ws.cell(r, col).value or '').strip().upper() == n: return r
    for r in range(r0, r1):
        if not ws.cell(r, col).value: return r
def grupo(prefijo):
    g = collections.defaultdict(dict)
    for k, v in f.items():
        if k.startswith(prefijo) and v:
            base, campo = k.rsplit('__', 1); g[base][campo] = v
    return g

# A. Telas
for base, d in sorted(grupo('TELA_').items()):
    if not d.get('nombre'): continue
    r = fila_de(P, 1, d['nombre'], 5, 125)
    poner(P, r, 1, d['nombre'].strip(), 'tela')
    if d.get('tipo'): poner(P, r, 2, d['tipo'].strip().capitalize(), f"tipo {d['nombre']}")
    if d.get('unidad'): poner(P, r, 4, 'metro' if d['unidad'].strip().upper().startswith('M') else 'kg', f"se compra en {d['nombre']}")
    poner(P, r, 5, num(d.get('precio')), f"precio actual {d['nombre']}")
    if d.get('iva'): poner(P, r, 6, 'Sí' if d['iva'].strip().upper().startswith('S') else 'No', f"IVA {d['nombre']}")
    poner(P, r, 7, num(d.get('kgm')), f"kilos por metro {d['nombre']}")
    poner(P, r, 8, num(d.get('ancho')), f"ancho útil {d['nombre']}")
# B. Impresión
for key, row in (('papel_rollo', 5), ('papel_metros', 6), ('tinta_litro', 7), ('tinta_ml', 8), ('proteccion', 9)):
    poner(P, row, 21, num(f.get(f'IMP__{key}')), key)
# C. Confección
for base, d in sorted(grupo('CONF_').items()):
    if d.get('tipo'): r = fila_de(P, 14, d['tipo'], 5, 85); poner(P, r, 14, d['tipo'].strip(), 'confección'); poner(P, r, 15, num(d.get('precio')), f"$ {d['tipo']}")
# D. Avíos
for base, d in sorted(grupo('AVIO_').items()):
    if d.get('nombre'): r = fila_de(P, 17, d['nombre'], 5, 65); poner(P, r, 17, d['nombre'].strip(), 'avío'); poner(P, r, 18, num(d.get('precio')), f"$ {d['nombre']}")
# E. Moldes y despiece
MARCADAS = {}
codigos = {san(M.cell(r, 1).value): (r, M.cell(r, 1).value) for r in range(5, 155) if M.cell(r, 1).value}
# Moldes nuevos (hojas de despiece): se agregan a Moldes y sus piezas quedan ligadas al número de la hoja
for base, d in grupo('MOLDENUEVO_').items():
    numero = base.split('_', 1)[1]; cod = (d.get('codigo') or '').strip()
    if not cod or cod.rstrip(' ·') == numero: continue
    if san(cod) in codigos: codigos[numero] = codigos[san(cod)]; continue
    r = fila_de(M, 1, cod, 5, 155)
    poner(M, r, 1, cod, 'molde nuevo'); poner(M, r, 2, d.get('descripcion'), f'descripción {cod}')
    poner(M, r, 3, d.get('prenda'), f'prenda {cod}'); poner(M, r, 4, d.get('calce'), f'calce {cod}')
    codigos[san(cod)] = codigos[numero] = (r, cod)
for base, d in grupo('MOLDE__').items():
    s = base.split('__', 1)[1]
    if s not in codigos: continue
    r, code = codigos[s]
    poner(M, r, 5, d.get('confeccion'), f'confección {code}')
    if d.get('avios'):
        partes = [re.match(r'\s*(.+?)\s*[xX*]\s*([0-9.,]+)\s*$', p) for p in d['avios'].split('+')]
        for (ca, cq), m in zip(((6, 7), (8, 9)), [p for p in partes if p]):
            poner(M, r, ca, m.group(1), f'avío {code}'); poner(M, r, cq, num(m.group(2)), f'cant. avío {code}')
    marc = [f"{k} {d[k]}" for k in ('marcada_tela', 'marcada_metros', 'marcada_prendas', 'aprovechamiento', 'talla', 'suma_piezas', 'obs') if d.get(k)]
    for n in (1, 2, 3):   # marcadas por tela (hojas de despiece)
        mm = [f"{k} {d[f'm{n}_{k}']}" for k in ('tela', 'ancho', 'metros', 'prendas', 'aprov') if d.get(f'm{n}_{k}')]
        if any(d.get(f'm{n}_{k}') for k in ('ancho', 'metros', 'prendas', 'aprov')): marc.append(f'marcada {n}: ' + ' '.join(mm))
        MARCADAS[(code, n)] = (d.get(f'm{n}_tela'), num(d.get(f'm{n}_ancho')))
    if any(d.get(k) for k in ('marcada_metros', 'marcada_prendas', 'aprovechamiento', 'suma_piezas', 'obs')) or any(x.startswith('marcada ') and ':' in x for x in marc):
        nota = 'Marcada ' + (d.get('fecha') or '') + ': ' + ', '.join(marc).replace('marcada_', '').replace('_', ' ')
        old = str(M.cell(r, 17).value or '')
        if nota not in old: poner(M, r, 17, (old + ' | ' if old else '') + nota, f'nota marcada {code}')
    aj = num(d.get('ajuste'))
    if aj is not None: poner(M, r, 11, aj / 100, f'ajuste real {code}')
    poner(M, r, 18, num(d.get('ancho_ref')), f'ancho de referencia {code}')
filas_molde = collections.defaultdict(list)
for r in range(5, 1005):
    if D.cell(r, 1).value: filas_molde[D.cell(r, 1).value].append(r)
libres = [r for r in range(5, 1005) if not D.cell(r, 1).value]
for base, d in sorted(grupo('PZA__').items()):
    mn = int(num(d.get('marcada')) or 0)
    _, s, j = base.split('__'); j = int(j)
    if s not in codigos or not d.get('pieza'): continue
    code = codigos[s][1]; fm = filas_molde[code]
    r = next((x for x in fm if str(D.cell(x, 2).value).strip().upper() == d['pieza'].strip().upper()), None)
    if r is None:
        r = libres.pop(0); poner(D, r, 1, code, 'molde nuevo en despiece'); fm.append(r)
    poner(D, r, 2, d['pieza'].strip(), f'pieza {code}'); poner(D, r, 3, num(d.get('cant')), f"cant. {code} {d['pieza']}")
    poner(D, r, 4, num(d.get('metros')), f"metros {code} {d['pieza']}"); poner(D, r, 5, d.get('tela'), f"tela {code} {d['pieza']}")
    mt, ma = MARCADAS.get((code, mn), (None, None))
    if mt and not d.get('tela'): poner(D, r, 5, mt.strip(), f"tela {code} {d['pieza']} (de la marcada {mn})")
    poner(D, r, 9, num(d.get('ancho')) or ma, f"ancho de la marcada {code} {d['pieza']}")
    if d.get('sublima'): poner(D, r, 6, 'Sí' if d['sublima'].strip().upper().startswith('S') else 'No', f"sublima {code} {d['pieza']}")
# F. Cortes reales -> ajuste por molde (si no se escribió a mano)
kgm = {str(P.cell(r, 1).value).strip().upper(): P.cell(r, 7).value for r in range(5, 125) if P.cell(r, 1).value}
ajustes = collections.defaultdict(list)
for base, d in sorted(grupo('CORTE_').items()):
    if not num(d.get('kg')) and num(d.get('kg_antes')) and num(d.get('kg_despues')) is not None:
        d['kg'] = str(num(d['kg_antes']) - num(d['kg_despues'])).replace('.', ',')
    if not (d.get('molde') and d.get('tela') and num(d.get('prendas')) and num(d.get('kg'))): continue
    code = next((c for _, c in codigos.values() if c.upper().startswith(d['molde'].strip().upper()) or d['molde'].strip().upper() in c.upper()), None)
    k = kgm.get(d['tela'].strip().upper())
    if not code or not k: cambios.append(f"CORTE no usado ({d}): molde o peso de 1 m de la tela no encontrado"); continue
    real = num(d['kg']) / num(d['prendas']) / k
    desp = sum(D.cell(r, 4).value or 0 for r in filas_molde[code] if str(D.cell(r, 5).value or '').strip().upper() == d['tela'].strip().upper())
    if desp: ajustes[code].append(real / desp - 1)
    nota = f"Corte {d.get('fecha','')} OT {d.get('ot','')}: {d['kg']} kg {d['tela']} / {d['prendas']} prendas = {real:.3f} m por prenda"
    r = codigos[san(code)][0]; old = str(M.cell(r, 17).value or '')
    if nota not in old: poner(M, r, 17, (old + ' | ' if old else '') + nota, f'nota corte {code}')
for code, lst in ajustes.items():
    r = codigos[san(code)][0]
    if not grupo('MOLDE__').get(f'MOLDE__{san(code)}', {}).get('ajuste'):
        poner(M, r, 11, round(sum(lst) / len(lst), 4), f'ajuste real {code} (promedio de {len(lst)} corte(s))')

wb.save(OUT)
print(f'{len(cambios)} cambios'); print('\n'.join(cambios))
