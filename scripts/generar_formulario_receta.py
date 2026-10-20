# PDF rellenable con los datos de la receta (precargado desde PMK_Receta_Produccion.xlsx).
# Los nombres de campo describen el dato, para poder leerlo de vuelta y actualizar el Excel.
import sys, re
from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, black, white

SRC, DST = sys.argv[1], sys.argv[2]
wb = load_workbook(SRC)
P, M, D = wb['Precios'], wb['Moldes'], wb['Despiece']
# Precios: telas A=tela B=tipo D=se compra en E=precio actual F=IVA G=kg/m H=ancho; confección N:O; avíos Q:R; parámetros U
telas = [(r, P.cell(r, 1).value, P.cell(r, 2).value, P.cell(r, 4).value, P.cell(r, 5).value, P.cell(r, 6).value, P.cell(r, 7).value, P.cell(r, 8).value) for r in range(5, 125) if P.cell(r, 1).value]
conf = [(r, P.cell(r, 14).value, P.cell(r, 15).value) for r in range(5, 85) if P.cell(r, 14).value]
avios = [(r, P.cell(r, 17).value, P.cell(r, 18).value) for r in range(5, 65) if P.cell(r, 17).value]
par = {k: P.cell(r, 21).value for k, r in (('papel_rollo', 5), ('papel_metros', 6), ('tinta_litro', 7), ('tinta_ml', 8), ('proteccion', 9), ('merma', 12), ('gg', 14), ('meta', 15))}
moldes = [(r, [M.cell(r, c).value for c in range(1, 19)]) for r in range(5, 155) if M.cell(r, 1).value]
piezas = {}
for r in range(5, 1005):
    if D.cell(r, 1).value: piezas.setdefault(D.cell(r, 1).value, []).append((r, D.cell(r, 2).value, D.cell(r, 3).value, D.cell(r, 4).value, D.cell(r, 5).value, D.cell(r, 6).value, D.cell(r, 9).value))

AZUL = HexColor('#1F3864'); GRIS = HexColor('#595959'); FONDO = HexColor('#EAF1FB'); LINEA = HexColor('#BFBFBF')
W, H = letter; X0, X1 = 36, W - 36
c = canvas.Canvas(DST, pagesize=letter); c.setTitle('PMK - Datos para la receta de costo de produccion')
form = c.acroForm
pag = [0]
def san(s): return re.sub(r'[^A-Za-z0-9]+', '_', str(s or '')).strip('_')[:40]
def num(v):
    if v is None or v == '': return ''
    if isinstance(v, float):
        s = f'{v:.4f}'.rstrip('0').rstrip('.')
        return s.replace('.', ',')
    return str(v)
def pie():
    c.setFont('Helvetica', 7); c.setFillColor(GRIS)
    c.drawString(X0, 20, 'Playmaker - Datos para la receta de costo. Completar y devolver el PDF (o escaneado) para actualizar PMK_Receta_Produccion.xlsx.')
    c.drawRightString(X1, 20, f'Pagina {pag[0]}')
def nueva(titulo, sub):
    if pag[0]: pie(); c.showPage()
    pag[0] += 1
    c.setFillColor(AZUL); c.rect(0, H - 50, W, 50, fill=1, stroke=0)
    c.setFillColor(white); c.setFont('Helvetica-Bold', 14); c.drawString(X0, H - 30, titulo)
    c.setFont('Helvetica', 8); c.drawString(X0, H - 43, sub)
    c.setFillColor(black)
    return H - 66
def campo(nombre, x, y, w, valor='', h=13, tip=''):
    form.textfield(name=nombre, tooltip=tip or nombre, x=x, y=y - 3, width=w, height=h, value=str(valor or ''),
                   fontName='Helvetica', fontSize=8, borderWidth=0.5, borderColor=LINEA, fillColor=FONDO, textColor=black, forceBorder=True)
def responsable(y, sec):
    c.setFont('Helvetica', 8); c.drawString(X0, y, 'Completado por:'); campo(f'{sec}__responsable', X0 + 70, y, 170)
    c.drawString(X0 + 260, y, 'Fecha:'); campo(f'{sec}__fecha', X0 + 292, y, 80)
    return y - 22
def encabezado(y, cols):
    c.setFillColor(HexColor('#D9D9D9')); c.rect(X0, y - 4, X1 - X0, 14, fill=1, stroke=0); c.setFillColor(black)
    c.setFont('Helvetica-Bold', 7.5)
    for x, t in cols: c.drawString(x, y, t)
    return y - 17
def texto(y, lineas, size=8, color=GRIS):
    c.setFont('Helvetica', size); c.setFillColor(color)
    for l in lineas: c.drawString(X0, y, l); y -= size + 3
    c.setFillColor(black); return y - 4

# ---------------------------------------------------------------- portada / instrucciones
y = nueva('DATOS PARA LA RECETA DE COSTO DE PRODUCCION', 'Playmaker - formulario para completar por area y devolver')
y = texto(y, ['Como usar este formulario:',
              '1. Cada area completa SOLO su seccion. Los casilleros celestes se pueden escribir en el computador o, impreso, a mano con letra clara.',
              '2. Los datos que ya existen vienen precargados: corrige solo lo que cambio y completa lo vacio.',
              '3. Numeros: decimales con coma (0,197). Montos sin puntos ni signo $ (5600). Todo NETO, sin IVA, salvo donde se pide con IVA.',
              '4. Guarda el PDF (o escanealo) y subelo al chat: con eso se actualiza el archivo PMK_Receta_Produccion.xlsx.',
              '',
              'Secciones:',
              '   A. Telas (Compras) ........................ como se compra (kg o metro), precio actual, IVA, KILOS POR METRO y ancho util',
              '   B. Impresion (Sublimacion) ................ metros del rollo de papel y ml de tinta por metro impreso',
              '   C. Confeccion (Talleres) .................. precio por prenda a trato (factura: sin IVA; boleta: monto completo)',
              '   D. Avios (Compras) ........................ precio neto por unidad',
              '   E. Despiece por molde (Diseno / ploteo) ... metros por prenda de cada pieza, desde la marcada del plotter',
              '   F. Cortes reales (Corte) .................. tela, prendas cortadas y kg usados en cada corte medido'], 8.5, black)

# ---------------------------------------------------------------- A. Telas
cols = [(X0 + 2, 'Tela'), (X0 + 155, 'Tipo'), (X0 + 208, 'Se compra'), (X0 + 256, 'Precio actual'), (X0 + 319, 'IVA'), (X0 + 352, 'Kg por metro'), (X0 + 410, 'Ancho util (m)')]
def tabla_telas(y):
    return encabezado(y, cols)
y = nueva('A. TELAS - Compras / administracion', 'Kg por metro: cortar 1 m lineal de orilla a orilla y pesarlo (3 decimales). Ancho util: sin los orillos que no se pueden cortar.')
y = responsable(y, 'A'); y = tabla_telas(y)
filas = [t for t in telas] + [(None, '', '', '', '', '', '', '')] * 8
for i, (r, n, tp, un, pr, iv, kg, an) in enumerate(filas):
    if y < 50: y = nueva('A. TELAS (continuacion)', ''); y = tabla_telas(y)
    k = f'TELA_{i+1:03d}'
    campo(f'{k}__nombre', X0, y, 152, n); campo(f'{k}__tipo', X0 + 154, y, 50, tp); campo(f'{k}__unidad', X0 + 206, y, 46, un)
    campo(f'{k}__precio', X0 + 254, y, 61, num(pr)); campo(f'{k}__iva', X0 + 317, y, 31, iv)
    campo(f'{k}__kgm', X0 + 350, y, 56, num(kg)); campo(f'{k}__ancho', X0 + 408, y, 56, num(an))
    y -= 16

# ---------------------------------------------------------------- B. Impresión
y = nueva('B. IMPRESION - Sublimacion', 'ml de tinta por metro = litros usados en el mes / metros impresos del mes (Control de Ploteo) x 1000.')
y = responsable(y, 'B')
for key, lab, v in (('papel_rollo', 'Papel sublimacion: precio del rollo, neto ($)', par['papel_rollo']), ('papel_metros', 'Papel sublimacion: metros reales que rinde el rollo', par['papel_metros']),
                    ('tinta_litro', 'Tinta: precio por litro, neto ($)', par['tinta_litro']), ('tinta_ml', 'Tinta: ml por metro impreso', par['tinta_ml']),
                    ('proteccion', 'Papel de proteccion (kraft): $ por metro (si se usa)', par['proteccion']),
                    ('litros_mes', 'Dato de respaldo: litros de tinta usados en el mes', ''), ('metros_mes', 'Dato de respaldo: metros impresos en el mes', ''), ('mes', 'Mes de la medicion', '')):
    c.setFont('Helvetica', 9); c.drawString(X0, y, lab); campo(f'IMP__{key}', X0 + 300, y, 120, num(v)); y -= 22

# ---------------------------------------------------------------- C. Confección
def enc_conf(y): return encabezado(y, [(X0 + 2, 'Tipo de confeccion'), (X0 + 300, '$ por prenda (taller)'), (X0 + 400, 'Taller / documento')])
y = nueva('C. CONFECCION - Talleres', 'Taller con factura (Paola): precio SIN IVA. Taller con boleta de honorarios (Sandra, Jacqueline, Sonia, Antonella): monto completo.')
y = responsable(y, 'C'); y = enc_conf(y)
for i, (r, n, v) in enumerate(conf + [(None, '', '')] * 6):
    if y < 50: y = nueva('C. CONFECCION (continuacion)', ''); y = enc_conf(y)
    k = f'CONF_{i+1:03d}'
    campo(f'{k}__tipo', X0, y, 290, n); campo(f'{k}__precio', X0 + 298, y, 90, num(v)); campo(f'{k}__taller', X0 + 398, y, 142, '')
    y -= 16

# ---------------------------------------------------------------- D. Avíos
y = nueva('D. AVIOS - Compras', 'Precio neto por unidad.')
y = responsable(y, 'D'); y = encabezado(y, [(X0 + 2, 'Avio'), (X0 + 260, '$ por unidad, neto'), (X0 + 360, 'Proveedor / nota')])
for i, (r, n, v) in enumerate(avios + [(None, '', '')] * 8):
    k = f'AVIO_{i+1:03d}'
    campo(f'{k}__nombre', X0, y, 250, n); campo(f'{k}__precio', X0 + 258, y, 90, num(v)); campo(f'{k}__nota', X0 + 358, y, 182, '')
    y -= 16

# ---------------------------------------------------------------- E. Despiece
SUB = 'Metros por prenda de cada fila (ya sumando la cantidad), talla base, desde la marcada de SU tela. Ancho marcada: ancho util de esa tela. Sublima: Si / No.'
def enc_desp(y): return encabezado(y, [(X0 + 2, 'Pieza'), (X0 + 150, 'Cant.'), (X0 + 182, 'METROS x prenda'), (X0 + 255, 'Ancho marcada'), (X0 + 320, 'Tela por defecto'), (X0 + 486, 'Sublima')])
y = nueva('E. DESPIECE POR MOLDE - Diseno / moldaje / ploteo', SUB); y = responsable(y, 'E')
for r, v in moldes:
    code = v[0]; pz = piezas.get(code, [])
    need = 50 + 16 * (len(pz) + 2)
    if y - need < 40: y = nueva('E. DESPIECE POR MOLDE (continuacion)', SUB)
    mk = f'MOLDE__{san(code)}'
    c.setFillColor(AZUL); c.setFont('Helvetica-Bold', 9.5); c.drawString(X0, y, f'{code}'); c.setFillColor(GRIS); c.setFont('Helvetica', 8)
    c.drawString(X0 + 150, y, f'{v[1] or ""}  -  {v[2] or ""} {v[3] or ""}'[:70]); c.setFillColor(black); y -= 16
    c.setFont('Helvetica', 7.5); c.drawString(X0, y, 'Confeccion:'); campo(f'{mk}__confeccion', X0 + 45, y, 165, v[4])
    c.drawString(X0 + 218, y, 'Avios (nombre x cant.):'); campo(f'{mk}__avios', X0 + 305, y, 120, ' + '.join(f'{a} x {q}' for a, q in ((v[5], v[6]), (v[7], v[8])) if a))
    c.drawString(X0 + 432, y, 'Ajuste real %:'); campo(f'{mk}__ajuste', X0 + 488, y, 52, num(v[10] * 100) if isinstance(v[10], (int, float)) else '')
    y -= 16; c.drawString(X0, y, 'Ancho de la tela con que se midio el despiece (m):'); campo(f'{mk}__ancho_ref', X0 + 190, y, 50, num(v[17]))
    y -= 18; y = enc_desp(y)
    for j, (dr, pzn, q, mt, tl, sb, an) in enumerate(pz + [(None, '', '', '', '', '', '')] * 2):
        k = f'PZA__{san(code)}__{j+1:02d}'
        campo(f'{k}__pieza', X0, y, 146, pzn); campo(f'{k}__cant', X0 + 148, y, 30, q); campo(f'{k}__metros', X0 + 180, y, 70, num(mt))
        campo(f'{k}__ancho', X0 + 253, y, 62, num(an)); campo(f'{k}__tela', X0 + 318, y, 164, tl); campo(f'{k}__sublima', X0 + 484, y, 40, sb)
        y -= 16
    y -= 8

# ---------------------------------------------------------------- F. Cortes reales
def enc_corte(y): return encabezado(y, [(X0 + 2, 'Fecha'), (X0 + 52, 'OT'), (X0 + 100, 'Molde (codigo)'), (X0 + 230, 'Tela cortada'), (X0 + 385, 'Prendas'), (X0 + 430, 'Kg usados'), (X0 + 482, 'Talla base')])
y = nueva('F. CORTES REALES - Corte', 'Pesar la tela antes de cortar y la que vuelve al rollo despues. Kg usados = antes - despues (los retazos cuentan como usados). Si hay 2 telas, una fila por tela.')
y = responsable(y, 'F'); y = enc_corte(y)
for i in range(30):
    if y < 50: y = nueva('F. CORTES REALES (continuacion)', ''); y = enc_corte(y)
    k = f'CORTE_{i+1:03d}'
    campo(f'{k}__fecha', X0, y, 48); campo(f'{k}__ot', X0 + 50, y, 46); campo(f'{k}__molde', X0 + 98, y, 128); campo(f'{k}__tela', X0 + 228, y, 152)
    campo(f'{k}__prendas', X0 + 383, y, 42); campo(f'{k}__kg', X0 + 428, y, 50); campo(f'{k}__talla', X0 + 480, y, 60)
    y -= 16
pie(); c.save()
print('ok', pag[0], 'paginas', len(telas), 'telas', len(moldes), 'moldes', sum(len(v) for v in piezas.values()), 'piezas')
